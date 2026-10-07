"""Combine the Fantasy Legends data files into one Excel workbook (for uploading to ChatGPT etc.).

Usage: python build_workbook.py ../data ../data/fantasy_legends.xlsx
"""

import json
import os
import sys

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

NOTES = [
    ("Fantasy Legends: NFL game logs, 1950-1999", ""),
    ("", ""),
    ("Sheet", "What's in it"),
    ("Players", "92 star QBs, RBs, WRs, TEs and kickers: position, Hall of Fame flag, seasons, teams, college, draft"),
    ("Player Game Logs", "One row per game per player (regular season + playoffs). Match to Players on player_id"),
    ("Team Defense Logs", "One row per team per game, every team 1950-1999"),
    ("Featured Defenses", "Regular-season totals for iconic defenses (Steel Curtain, '85 Bears, etc.)"),
    ("", ""),
    ("Column", "Meaning"),
    ("week", "The team's game number in that season"),
    ("playoff", "TRUE = postseason game"),
    ("home_away", "H = home, A = away, N = neutral site"),
    ("result", "W / L / T for the player's (or defense's) team"),
    ("age", "Player age on game day, years-days (28-284 = 28 years, 284 days)"),
    ("sacked", "Times the QB was sacked"),
    ("ret_td", "Kick and punt return touchdowns"),
    ("xpm / xpa, fgm / fga", "Extra points and field goals made / attempted"),
    ("Team codes", "Pro-Football-Reference city codes, so a code follows the city, not the franchise. "
                   "BAL = Colts 1953-83 and Ravens 1996-99; IND = Colts 1984+; STL = Cardinals 1960-87 and Rams 1995-99; "
                   "CRD = Chicago Cardinals; RAM = LA Rams; RAI = LA Raiders 1982-94; OAK = Raiders; "
                   "PHO / ARI = Cardinals; HOU = Oilers; TEN = Titans; BOS / NWE = Patriots; "
                   "NYT = NY Titans (Jets); DTX = Dallas Texans; LAC / SDG = Chargers; GNB = Green Bay; "
                   "KAN, NOR, TAM, SFO = Kansas City, New Orleans, Tampa Bay, San Francisco"),
    ("", ""),
    ("Fantasy scoring", ""),
    ("fpts_std", "Pass: 0.04/yd, 4/TD, -2/INT. Rush and receiving: 0.1/yd, 6/TD. Return TD: 6. Kicker: 3/FG, 1/XP"),
    ("fpts_ppr", "Same as fpts_std plus 1 point per reception"),
    ("Team defense fpts", "1/sack, 2/INT, 2/safety, 6/INT-return or return TD, plus points allowed: "
                          "0 = 10, 1-6 = 7, 7-13 = 4, 14-20 = 1, 21-27 = 0, 28-34 = -1, 35+ = -4"),
    ("", ""),
    ("Gaps in the data", ""),
    ("Fumbles, 2-pt conversions, FG distance", "Not in the source data, so not counted in fantasy points"),
    ("Sacks", "Not an official stat before 1982: blank for team defenses before then, QB 'sacked' often 0"),
    ("Targets", "Blank before 1992"),
    ("Team defense", "No fumble recoveries or fumble-return TDs"),
    ("Season length", "12 games through 1960 (14 for the 1960 AFL), 14 for 1961-1977, 16 from 1978. "
                      "Strike years: 9 games in 1982, 15 in 1987. Compare eras per game, not season totals"),
    ("Careers past 1999", "Only games through 1999 are included (Rice, Favre, etc.)"),
    ("", ""),
    ("Source", "Pro-Football-Reference via Kaggle dataset zynicide/nfl-football-player-stats (scraped Dec 2017). "
               "Pass completions/attempts were swapped in the source and are corrected here"),
]


def load_columnar(path):
    d = json.load(open(path))
    return pd.DataFrame(d["rows"], columns=d["columns"])


def main(data_dir, out_path):
    players = pd.DataFrame(json.load(open(os.path.join(data_dir, "players.json"))))
    players["teams"] = players["teams"].str.join(", ")
    logs = load_columnar(os.path.join(data_dir, "player_gamelogs.json"))
    dst = load_columnar(os.path.join(data_dir, "dst_gamelogs.json"))
    featured = pd.DataFrame(json.load(open(os.path.join(data_dir, "featured_defenses.json"))))

    sheets = {"Players": players, "Player Game Logs": logs,
              "Team Defense Logs": dst, "Featured Defenses": featured}
    with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
        pd.DataFrame(NOTES).to_excel(xw, sheet_name="Notes", index=False, header=False)
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)

        font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
        head_fill = PatternFill("solid", start_color="D9E1F2")
        for ws in xw.book.worksheets:
            name = ws.title
            for row in ws.iter_rows():
                for cell in row:
                    cell.font = font
            if name == "Notes":
                ws.column_dimensions["A"].width = 38
                ws.column_dimensions["B"].width = 120
                ws["A1"].font = Font(name="Arial", size=14, bold=True)
                for row in ws.iter_rows(min_col=1, max_col=2):
                    if row[0].value in ("Sheet", "Column", "Fantasy scoring", "Gaps in the data", "Source"):
                        for c in row:
                            c.font = bold
                continue
            for cell in ws[1]:
                cell.font, cell.fill = bold, head_fill
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(sheets[name].columns, 1):
                width = max(len(str(col)), sheets[name][col].astype(str).str.len().head(500).max())
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 6), 40)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
